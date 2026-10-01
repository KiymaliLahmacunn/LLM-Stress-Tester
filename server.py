"""
server.py
---------
FastAPI application: REST API endpoints for the LLM Stress Test platform.
"""

import json
import os
import urllib.request
from urllib.error import URLError

import openpyxl
import requests
from dotenv import load_dotenv
from fastapi import FastAPI, File, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from core.dataset_manager import (
    delete_dataset,
    process_and_save_upload,
)
from core.dataset_manager import (
    list_datasets as get_datasets,
)
from core.math_report import generate_report
from core.test_manager import TestManager
from core.utils import OUR_SERVER_PORT, TARGET_PORTS, detect_hardware

load_dotenv()

app = FastAPI(title="LLM Stress Tester")
manager = TestManager()

os.makedirs("static", exist_ok=True)
os.makedirs("templates", exist_ok=True)

app.mount("/static", StaticFiles(directory="static"), name="static")

DATASETS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "datasets")
os.makedirs(DATASETS_DIR, exist_ok=True)


# ── Pages ────────────────────────────────────────────────────────────────


@app.get("/favicon.ico", include_in_schema=False)
def favicon():
    return Response(content=b"", media_type="image/x-icon")


@app.get("/", response_class=HTMLResponse)
async def get_index():
    with open("templates/index.html", "r", encoding="utf-8") as f:
        return f.read()


# ── Model Discovery ─────────────────────────────────────────────────────


@app.get("/api/models")
def get_models():
    hardware_info, hw_color = detect_hardware()
    found = []

    for server_name, port in TARGET_PORTS.items():
        if port == OUR_SERVER_PORT:
            continue

        try:
            url = f"http://127.0.0.1:{port}/v1/models"
            resp = requests.get(url, timeout=1.5)
            if resp.status_code != 200:
                continue

            models = resp.json().get("data", [])
            for m in models:
                raw_name = m.get("id", "Unknown")
                parts = raw_name.split("-")
                display_name = parts[0].capitalize() if parts else raw_name
                version = "-".join(parts[1:])[:15] if len(parts) > 1 else "Default"

                found.append(
                    {
                        "name": display_name,
                        "raw_name": raw_name,
                        "version": version,
                        "location": f"{server_name}:{port}",
                        "hardware": hardware_info,
                        "color": hw_color,
                        "full_api_url": f"http://127.0.0.1:{port}/v1/chat/completions",
                    }
                )
        except Exception:
            continue

    return {"models": found}


# ── Datasets ─────────────────────────────────────────────────────────────


@app.get("/api/datasets")
def list_datasets_endpoint():
    return {"datasets": get_datasets()}


@app.post("/api/datasets/upload")
async def upload_dataset(file: UploadFile = File(...)):
    try:
        content = await file.read()
        if len(content) > 10 * 1024 * 1024:
            return {"error": "File size exceeds 10 MB limit."}
        metadata = process_and_save_upload(content, file.filename)
        return {"status": "uploaded", "filename": metadata["name"]}
    except Exception as e:
        return {"error": str(e)}


@app.get("/api/datasets/download/{filename}")
def download_dataset(filename: str):
    safe_name = _safe_filename(filename)
    if not safe_name:
        return {"error": "invalid filename"}

    path = os.path.join(DATASETS_DIR, safe_name)
    if os.path.exists(path):
        return FileResponse(path, filename=safe_name)
    return {"error": "not found"}


@app.delete("/api/datasets/delete/{filename}")
def delete_dataset_endpoint(filename: str):
    try:
        if manager.is_testing and manager.active_dataset_path:
            active_basename = os.path.basename(manager.active_dataset_path)
            if active_basename == filename:
                return {
                    "error": "Cannot delete dataset while it is being used in an active test."
                }
        delete_dataset(filename)
        return {"status": "deleted"}
    except Exception as e:
        return {"error": str(e)}


@app.get("/api/datasets/preview/{filename}")
def preview_dataset(filename: str):
    try:
        from core.dataset_manager import get_dataset_preview
        preview = get_dataset_preview(filename)
        return {"preview": preview}
    except Exception as e:
        return {"error": str(e)}


# ── Helpers ──────────────────────────────────────────────────────────────


def _safe_filename(filename: str) -> str | None:
    """
    Validate that filename has no path traversal components.
    Returns the sanitized basename or None if unsafe.
    """
    basename = os.path.basename(filename)
    if not basename or basename != filename or ".." in filename:
        return None
    return basename


# ── Test Control ─────────────────────────────────────────────────────────


class StartTestRequest(BaseModel):
    provider: str = "openai"
    model_name: str
    model_url: str
    dataset_filename: str
    concurrency: int
    rps: int
    duration: int
    thinking: bool
    streaming: bool
    hardware: str = ""
    seed: int | None = None
    timeout: int = 30


def _find_file_in_results(safe_name: str):
    base_dir = os.path.dirname(os.path.abspath(__file__))
    results_dir = os.path.join(base_dir, "results")
    if not os.path.exists(results_dir):
        return None
    for root, dirs, files in os.walk(results_dir):
        if safe_name in files:
            return os.path.join(root, safe_name)
    return None


@app.post("/api/test/start")
def start_test(req: StartTestRequest):
    dataset_path = os.path.join(DATASETS_DIR, req.dataset_filename)
    if not os.path.exists(dataset_path):
        return {"error": "Dataset not found"}

    # Auto-correct URL if user just entered base IP
    final_url = req.model_url.strip()
    if not final_url.startswith("http"):
        final_url = "http://" + final_url

    if not any(x in final_url for x in ["/chat", "/completions", "/generate", "/messages"]):
        if final_url.endswith("/v1") or final_url.endswith("/v1/"):
            final_url = final_url.rstrip("/") + "/chat/completions"
        else:
            final_url = final_url.rstrip("/") + "/v1/chat/completions"

    manager.configure_test(
        req.provider,
        req.model_name,
        final_url,
        dataset_path,
        req.concurrency,
        req.rps,
        req.duration,
        req.thinking,
        req.streaming,
        req.hardware,
        req.seed,
        req.timeout,
    )

    if not manager.active_dataset:
        return {
            "error": "Dataset is empty (0 questions). Please upload a dataset with at least one question."
        }

    manager.start_test()

    return {"status": "started"}


@app.post("/api/test/stop")
def stop_test():
    manager.stop_test()
    return {"status": "stopped"}


@app.get("/api/test/status")
def get_status(last_log_idx: int = 0):
    return manager.get_status(last_log_idx)


@app.get("/api/test/logs")
def get_logs():
    """Return a thread-safe snapshot of the in-memory logs."""
    return {"logs": manager.get_logs_snapshot()}


# ── Results ──────────────────────────────────────────────────────────────


@app.get("/api/results")
def list_results():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    results_dir = os.path.join(base_dir, "results")
    files = []
    seen_names = set()

    if os.path.exists(results_dir):
        for root, dirs, filenames in os.walk(results_dir):
            for name in filenames:
                if (
                    name.endswith(".xlsx")
                    and not name.startswith("~$")
                    and not name.endswith("_RUNNING.xlsx")
                ):
                    if name in seen_names:
                        continue
                    seen_names.add(name)

                    path = os.path.join(root, name)
                    mtime = os.path.getmtime(path)

                    meta = {}
                    meta_path = os.path.join(root, name.replace(".xlsx", ".json"))
                    if os.path.exists(meta_path):
                        try:
                            with open(meta_path, "r", encoding="utf-8") as mf:
                                meta = json.load(mf)
                        except json.JSONDecodeError, OSError:
                            pass

                    files.append({"name": name, "mtime": mtime, "meta": meta})

    files.sort(key=lambda x: x["mtime"], reverse=True)
    return {"files": files}


@app.get("/api/results/download/{filename}")
def download_result(filename: str):
    safe = _safe_filename(filename)
    if not safe:
        return {"error": "invalid filename"}

    path = _find_file_in_results(safe)
    if path and os.path.exists(path):
        return FileResponse(path, filename=safe)
    return {"error": "not found"}


@app.get("/api/results/report/{filename}")
def get_report(filename: str):
    safe = _safe_filename(filename)
    if not safe:
        return {"error": "invalid filename"}

    path = _find_file_in_results(safe)
    if path and os.path.exists(path):
        return generate_report(path)
    return {"error": "file not found"}


# ── Compare ──────────────────────────────────────────────────────────────


class CompareRequest(BaseModel):
    files: list[str]


@app.post("/api/compare")
def compare_results(req: CompareRequest):
    results = []

    for fname in req.files:
        safe = _safe_filename(fname)
        if not safe:
            continue

        path = _find_file_in_results(safe)
        if not path or not os.path.exists(path):
            continue

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                if "Summary" in wb.sheetnames:
                    ws = wb["Summary"]
                    headers = [
                        str(c.value) if c.value is not None else "" for c in ws[1]
                    ]
                    rows = []
                    for row in ws.iter_rows(min_row=2):
                        row_data = [
                            str(c.value) if c.value is not None else "" for c in row
                        ]
                        if any(v.strip() for v in row_data):
                            rows.append(row_data)
                    results.append(
                        {
                            "filename": safe,
                            "headers": headers,
                            "rows": rows,
                        }
                    )
            finally:
                wb.close()
        except Exception:
            pass

    return {"comparisons": results}


class DeleteResultsRequest(BaseModel):
    files: list[str]


@app.post("/api/results/delete")
def delete_results(req: DeleteResultsRequest):
    base_dir = os.path.dirname(os.path.abspath(__file__))
    results_dir = os.path.join(base_dir, "results")

    if not os.path.exists(results_dir):
        return {"status": "deleted", "deleted": 0}

    # Build a map of filename -> parent_dir in a single pass
    file_map = {}
    for root, dirs, files in os.walk(results_dir):
        for name in files:
            if name.endswith(".xlsx"):
                file_map[name] = root

    deleted_count = 0
    errors = []
    deleted_dirs = set()

    for fname in req.files:
        safe = _safe_filename(fname)
        if not safe or safe not in file_map:
            continue

        parent_dir = file_map[safe]

        # Skip if already deleted in this batch
        if parent_dir in deleted_dirs:
            deleted_count += 1
            continue

        if os.path.abspath(parent_dir).startswith(
            os.path.abspath(results_dir)
        ) and os.path.abspath(parent_dir) != os.path.abspath(results_dir):
            try:
                import shutil

                if os.path.exists(parent_dir):
                    shutil.rmtree(parent_dir)
                deleted_dirs.add(parent_dir)
                deleted_count += 1
            except Exception as e:
                errors.append(f"Failed to delete {safe}: {str(e)}")

    if errors:
        return {"error": "; ".join(errors), "deleted": deleted_count}
    return {"status": "deleted", "deleted": deleted_count}


# ── System ───────────────────────────────────────────────────────────────


@app.get("/api/system/hardware")
def get_hw():
    hw, _color = detect_hardware()
    return {"hardware": hw}


@app.get("/api/system/test_connection")
def test_connection(url: str):
    """
    1) Fetches /models to discover the model name.
    2) Sends a tiny test prompt to /chat/completions to verify inference works.
    3) Validates the model is text-to-text.
    Returns model name + prompt test result + text model flag.
    """
    try:
        url = url.strip()
        if not url.startswith("http://") and not url.startswith("https://"):
            url = "http://" + url

        if "/chat/completions" in url:
            base_url = url.split("/chat/completions")[0]
        elif "/completions" in url:
            base_url = url.split("/completions")[0]
        else:
            base_url = url.rstrip("/")
            if not base_url.endswith("/v1"):
                base_url += "/v1"

        models_url = f"{base_url}/models"
        chat_url = f"{base_url}/chat/completions"

        proxy_handler = urllib.request.ProxyHandler({})
        opener = urllib.request.build_opener(proxy_handler)

        # ── Step 1: Fetch model list ──
        req = urllib.request.Request(models_url, method="GET")
        req.add_header("Accept", "application/json")

        with opener.open(req, timeout=10.0) as response:
            data = json.loads(response.read().decode())

        first_model = ""
        if isinstance(data, dict):
            if (
                "data" in data
                and isinstance(data["data"], list)
                and len(data["data"]) > 0
            ):
                first_model = data["data"][0].get("id", "")
            elif (
                "models" in data
                and isinstance(data["models"], list)
                and len(data["models"]) > 0
            ):
                first_model = data["models"][0].get(
                    "name", data["models"][0].get("id", "")
                )
        elif isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
            first_model = data[0].get("name", "")

        if not first_model:
            first_model = "Unknown"

        # ── Step 2: Send a tiny test prompt ──
        prompt_ok = False
        prompt_reply = ""
        prompt_error = ""
        try:
            test_payload = json.dumps(
                {
                    "model": first_model,
                    "messages": [
                        {
                            "role": "system",
                            "content": "Provide a direct answer without any <think> tags, reasoning, or inner monologue.",
                        },
                        {"role": "user", "content": "Say hello in one word."},
                    ],
                    "max_tokens": 20,
                    "temperature": 0,
                }
            ).encode("utf-8")

            prompt_req = urllib.request.Request(
                chat_url, data=test_payload, method="POST"
            )
            prompt_req.add_header("Content-Type", "application/json")
            prompt_req.add_header("Accept", "application/json")

            with opener.open(prompt_req, timeout=60.0) as prompt_resp:
                prompt_data = json.loads(prompt_resp.read().decode())

            # Check if we got a valid response
            choices = prompt_data.get("choices", [])
            if choices and len(choices) > 0:
                msg = choices[0].get("message", {})
                content_str = msg.get("content", "") or ""
                reasoning_str = msg.get("reasoning_content", "") or ""
                prompt_reply = (content_str + reasoning_str)[:100]

                if prompt_reply:
                    prompt_ok = True
            elif "error" in prompt_data:
                prompt_error = str(prompt_data["error"])[:200]
        except urllib.error.HTTPError as he:
            try:
                err_body = json.loads(he.read().decode())
                prompt_error = str(err_body.get("error", {}).get("message", err_body))[
                    :200
                ]
            except Exception:
                prompt_error = f"HTTP {he.code}: {he.reason}"
        except Exception as pe:
            prompt_error = str(pe)[:200]

        # ── Step 3: Validate text-to-text model ──
        is_text_model = False
        non_text_patterns = [
            "embed",
            "embedding",
            "tts",
            "whisper",
            "dall-e",
            "stable-diffusion",
            "sdxl",
            "image",
            "vision-only",
            "audio",
            "speech",
            "nomic-embed",
            "text-embedding",
            "clip",
            "vocoder",
            "wav2vec",
        ]
        model_lower = first_model.lower()
        is_non_text = any(p in model_lower for p in non_text_patterns)

        if prompt_ok and not is_non_text:
            printable_ratio = sum(1 for c in prompt_reply if c.isprintable()) / max(
                len(prompt_reply), 1
            )
            is_text_model = printable_ratio > 0.8

        return {
            "status": "ok",
            "model": first_model,
            "prompt_ok": prompt_ok,
            "prompt_reply": prompt_reply,
            "prompt_error": prompt_error,
            "is_text_model": is_text_model,
        }

    except URLError as e:
        return {"status": "error", "error": f"Connection Error: {str(e.reason)}"}
    except Exception as e:
        return {"status": "error", "error": str(e)}
